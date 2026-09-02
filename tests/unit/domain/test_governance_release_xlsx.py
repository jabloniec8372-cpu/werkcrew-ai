"""Structural and semantic QA for the governance release workbook."""

from pathlib import Path
import re
from zipfile import ZipFile

import openpyxl
import pytest


WORKBOOK = (
    Path(__file__).parents[3]
    / "docs"
    / "governance"
    / "releases"
    / "v1.1"
    / "WERKcrew_State_Transition_Matrix_v1.1.xlsx"
)
SHEETS = [
    "00 Release",
    "01 Axes",
    "02 Promise",
    "03 Health Recovery",
    "04 Objections",
    "05 Voice Execution",
    "06 Regression 1017",
]


@pytest.fixture(scope="module")
def workbook():
    loaded = openpyxl.load_workbook(WORKBOOK, data_only=False)
    yield loaded
    loaded.close()


def test_release_workbook_is_complete_ooxml_with_frozen_topology(workbook) -> None:
    assert workbook.sheetnames == SHEETS
    assert workbook.properties.title == "WERKcrew State & Transition Matrix v1.1"
    assert workbook.calculation.calcMode in (None, "auto")
    assert workbook.calculation.fullCalcOnLoad

    with ZipFile(WORKBOOK) as package:
        entries = set(package.namelist())
    assert {
        "[Content_Types].xml",
        "xl/workbook.xml",
        "xl/styles.xml",
        "docProps/core.xml",
    } <= entries
    assert all(
        f"xl/worksheets/sheet{index}.xml" in entries
        for index in range(1, len(SHEETS) + 1)
    )

    with ZipFile(WORKBOOK) as package:
        workbook_xml = package.read("xl/workbook.xml").decode("utf-8")
    print_area_ids = {
        int(match)
        for match in re.findall(
            r'<definedName name="(?:_xlnm\.)?Print_Area" localSheetId="(\d+)"',
            workbook_xml,
        )
    }
    assert print_area_ids == set(range(len(SHEETS)))
@pytest.mark.parametrize("sheet_name", SHEETS)
def test_every_sheet_is_publication_ready(workbook, sheet_name: str) -> None:
    sheet = workbook[sheet_name]

    assert sheet.freeze_panes == "A5"
    assert sheet.page_setup.orientation == "landscape"
    if sheet_name != "00 Release":
        assert sheet.page_setup.fitToWidth == 1 or (
            sheet.page_setup.scale is not None and 10 <= sheet.page_setup.scale <= 100
        )
    assert sheet.page_setup.fitToHeight == 1
    assert sheet.sheet_view.showGridLines is False
    assert sheet["A1"].value
    assert sheet["A1"].font.bold
    assert sheet["A1"].fill.fill_type == "solid"
    assert not any(
        isinstance(cell.value, str)
        and any(marker in cell.value for marker in ("TODO", "TBD", "Lorem ipsum"))
        for row in sheet.iter_rows()
        for cell in row
    )


def test_promise_matrix_freezes_forced_and_hard_boundaries(workbook) -> None:
    sheet = workbook["02 Promise"]
    rows = {
        sheet.cell(row, 1).value: tuple(sheet.cell(row, column).value for column in range(2, 12))
        for row in range(5, 19)
    }

    assert rows["Forced soft exception"] == (
        "FEASIBLE",
        "SOFT_EXCEPTION",
        "APPROVED",
        "SUFFICIENT",
        "CLEAR",
        "VERIFIED",
        "NOT_REQUIRED",
        "FORCED",
        "ALLOWED",
        "EXECUTE_WITH_RECORDED_RISK",
    )
    assert rows["Unknown feasibility"][-2:] == (
        "BLOCKED",
        "ABSTAIN · unknown cannot become FORCED",
    )
    assert rows["BHP hard block"][-2:] == ("BLOCKED", "FORBIDDEN")
    assert rows["Missing required permission"][-2:] == ("BLOCKED", "FORBIDDEN")
    assert rows["Critical data conflicting"][-2:] == ("BLOCKED", "ABSTAIN")


def test_health_recovery_and_voice_records_remain_separate(workbook) -> None:
    health = workbook["03 Health Recovery"]
    assert [health.cell(14, column).value for column in range(1, 6)] == [
        "UNSTABLE",
        "RECOVERY_NEEDED",
        "DECLINED",
        "INACTIVE",
        "GROWTH",
    ]

    voice = workbook["05 Voice Execution"]
    assert [voice.cell(row, 1).value for row in range(13, 17)] == [
        "DRAFT_TEXT",
        "APPROVED_TEXT",
        "SENT_MESSAGE",
        "ACTUAL_CLAIM",
    ]
    assert [voice.cell(row, 1).value for row in range(5, 10)] == [
        "EXECUTE",
        "EXECUTE_WITH_RECORDED_RISK",
        "ABSTAIN",
        "FORBIDDEN",
        "HALT",
    ]


def test_regression_sheet_contains_all_blocking_cases(workbook) -> None:
    sheet = workbook["06 Regression 1017"]
    cases = {sheet.cell(row, 1).value: sheet.cell(row, 3).value for row in range(5, 20)}

    assert len(cases) == 15
    assert cases["T-PREF-01"] == "No veto; T1 allowed with recorded reserve risk"
    assert cases["T-AVAIL-01"] == "T2 = ABSTAIN"
    assert cases["T-METHOD-01"] == "Concrete T3 method = ABSTAIN"
    assert cases["T-BHP-01"] == "T3 = FORBIDDEN"
    assert cases["T-OUTCOME-01"] == "Does not rewrite the 10:17 assessment"
    assert cases["T-ROLE-01"] == "Same evidence gets the same content classification"


def test_release_dashboard_contains_live_formulas(workbook) -> None:
    sheet = workbook["00 Release"]
    assert [sheet.cell(row, 2).value for row in range(14, 20)] == [
        "docs/governance/WERKcrew_State_Transition_Matrix_v1.1.md"
    ] * 6
    assert sheet["B23"].value == "=COUNTA('01 Axes'!$C$5:$C$29)"
    assert sheet["B24"].value == "=COUNTA('06 Regression 1017'!$A$5:$A$19)"
    assert sheet["B25"].value == "=COUNTA('05 Voice Execution'!$A$5:$A$9)"
    assert [sheet.cell(row, 3).value for row in range(23, 26)] == [25, 15, 5]
