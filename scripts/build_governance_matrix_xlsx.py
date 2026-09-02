"""Build the publication XLSX projection of the governance matrix.

Markdown remains canonical. This builder intentionally contains presentation
data only and points every worksheet back to the versioned Markdown source.
"""

from __future__ import annotations

import argparse
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


NAVY = "17365D"
BLUE = "2E75B6"
PALE_BLUE = "D9EAF7"
PALE_GREEN = "E2F0D9"
PALE_AMBER = "FFF2CC"
PALE_RED = "F4CCCC"
PALE_PURPLE = "E4DFEC"
PALE_GRAY = "E7E6E6"
WHITE = "FFFFFF"
TEXT = "1F1F1F"
GRID = "B7C9D6"
SOURCE = "docs/governance/WERKcrew_State_Transition_Matrix_v1.1.md"


def _table(ws, ref: str, name: str) -> None:
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def _sheet_base(ws, title: str, subtitle: str, last_column: str) -> None:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A5"
    ws.merge_cells(f"A1:{last_column}1")
    ws["A1"] = title
    ws["A1"].font = Font(name="Calibri", size=18, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 30
    ws.merge_cells(f"A2:{last_column}2")
    ws["A2"] = subtitle
    ws["A2"].font = Font(name="Calibri", size=10, italic=True, color="44546A")
    ws["A2"].fill = PatternFill("solid", fgColor=PALE_BLUE)
    ws["A2"].alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 28
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.outlinePr.summaryBelow = True
    ws.oddHeader.center.text = "&BWERKcrew Decision Semantics Freeze"
    ws.oddFooter.left.text = "Source: State & Transition Matrix v1.1"
    ws.oddFooter.right.text = "Page &P of &N"
    ws.print_options.horizontalCentered = False


def _style_header(ws, row: int, start: int, end: int) -> None:
    for cell in ws.iter_cols(min_col=start, max_col=end, min_row=row, max_row=row):
        target = cell[0]
        target.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
        target.fill = PatternFill("solid", fgColor=BLUE)
        target.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        target.border = Border(bottom=Side(style="medium", color=NAVY))
    ws.row_dimensions[row].height = 30


def _style_body(ws, min_row: int, max_row: int, min_col: int, max_col: int) -> None:
    for row in ws.iter_rows(
        min_row=min_row,
        max_row=max_row,
        min_col=min_col,
        max_col=max_col,
    ):
        for cell in row:
            cell.font = Font(name="Calibri", size=10, color=TEXT)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = Border(bottom=Side(style="thin", color=GRID))


def _section_label(ws, row: int, text: str, last_column: str) -> None:
    ws.merge_cells(f"A{row}:{last_column}{row}")
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(name="Calibri", size=11, bold=True, color=NAVY)
    cell.fill = PatternFill("solid", fgColor=PALE_BLUE)
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 23


def _status_colors(ws, column: str, start: int, end: int) -> None:
    rules = [
        ("EXECUTE", PALE_GREEN),
        ("EXECUTE_WITH_RECORDED_RISK", PALE_AMBER),
        ("ABSTAIN", PALE_AMBER),
        ("FORBIDDEN", PALE_RED),
        ("HALT", PALE_PURPLE),
        ("ALLOWED", PALE_GREEN),
        ("BLOCKED", PALE_RED),
        ("FORCED", PALE_AMBER),
    ]
    target = f"{column}{start}:{column}{end}"
    for value, color in rules:
        ws.conditional_formatting.add(
            target,
            FormulaRule(
                formula=[f'ISNUMBER(SEARCH("{value}",{column}{start}))'],
                fill=PatternFill("solid", fgColor=color),
            ),
        )


def _overview(wb: Workbook) -> None:
    ws = wb.create_sheet("00 Release")
    ws.sheet_properties.tabColor = NAVY
    _sheet_base(
        ws,
        "WERKcrew State & Transition Matrix v1.1",
        "Publication workbook · Decision Semantics Freeze · Markdown is canonical",
        "H",
    )
    metadata = [
        ("Version", "1.1"),
        ("Status", "USTALONE — implementation contract"),
        ("Release date", date(2026, 9, 2)),
        ("Canonical source", SOURCE),
        ("ADR", "docs/decisions/0010-decision-semantics-freeze.md"),
    ]
    ws["A4"] = "Release metadata"
    ws["A4"].font = Font(bold=True, color=WHITE)
    ws["A4"].fill = PatternFill("solid", fgColor=BLUE)
    ws["B4"] = "Value"
    ws["B4"].font = Font(bold=True, color=WHITE)
    ws["B4"].fill = PatternFill("solid", fgColor=BLUE)
    for row, (label, value) in enumerate(metadata, start=5):
        ws.cell(row, 1, label)
        ws.cell(row, 2, value)
    ws["B7"].number_format = "yyyy-mm-dd"
    _style_body(ws, 5, 9, 1, 2)

    _section_label(ws, 11, "Workbook coverage", "H")
    ws.append([])
    coverage = [
        ("Axes and states", SOURCE, "A–F objects and frozen state vocabulary"),
        ("Promise gates", SOURCE, "FORCED derivation and hard-boundary behavior"),
        ("Health & recovery", SOURCE, "Observed health separated from OWNER controls"),
        ("Objections & consent", SOURCE, "Content-based classification and consent gate"),
        ("Voice & execution", SOURCE, "NO_DECEPTION and execution outcomes"),
        ("Regression contract", SOURCE, "DOM/KLINIKA/LOFT blocking scenarios"),
    ]
    ws.append(["Worksheet", "Source marker", "Purpose"])
    for item in coverage:
        ws.append(item)
    _style_header(ws, 13, 1, 3)
    _style_body(ws, 14, 19, 1, 3)
    _table(ws, "A13:C19", "ReleaseCoverage")

    _section_label(ws, 21, "Live checks", "H")
    checks = [
        ("Frozen axis states", "=COUNTA('01 Axes'!$C$5:$C$29)", 25),
        ("Regression cases", "=COUNTA('06 Regression 1017'!$A$5:$A$19)", 15),
        ("Execution outcomes", "=COUNTA('05 Voice Execution'!$A$5:$A$9)", 5),
    ]
    ws.append(["Check", "Formula result", "Expected"])
    for item in checks:
        ws.append(item)
    _style_header(ws, 22, 1, 3)
    _style_body(ws, 23, 25, 1, 3)
    _table(ws, "A22:C25", "ReleaseChecks")
    ws["B23"].number_format = "0"
    ws["B24"].number_format = "0"
    ws["B25"].number_format = "0"
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 52
    ws.column_dimensions["C"].width = 46
    for column in "DEFGH":
        ws.column_dimensions[column].width = 3
    ws.print_area = "A1:H25"


def _axes(wb: Workbook) -> None:
    ws = wb.create_sheet("01 Axes")
    ws.sheet_properties.tabColor = BLUE
    _sheet_base(ws, "Axes A–F", SOURCE, "D")
    rows = [
        ("A", "Declaration / speech act", "FIELD_TALK", "Field statement not yet faithfully recorded"),
        ("A", "Declaration / speech act", "UNCONFIRMED_BINDING_CLAIM", "Credible indication of a possible commitment; scope not confirmed"),
        ("A", "Declaration / speech act", "INTENT", "Intent without a binding date or scope"),
        ("A", "Declaration / speech act", "WINDOW", "Communicated range or explicit condition"),
        ("A", "Declaration / speech act", "PROMISE", "Binding commitment actually communicated"),
        ("B", "Concrete promise", "SAFE", "FEASIBLE + WITHIN_ENVELOPE + sufficient critical data"),
        ("B", "Concrete promise", "CONDITIONAL", "A named feasible condition remains outstanding"),
        ("B", "Concrete promise", "INFEASIBLE", "No physical or legal plan"),
        ("B", "Concrete promise", "FORCED", "Derived only: FEASIBLE + SOFT_EXCEPTION + OWNER APPROVED"),
        ("C", "Day or plan package", "FRAGILE", "Works only without a typical disruption"),
        ("C", "Day or plan package", "STABLE", "Absorbs defined typical deviations"),
        ("C", "Day or plan package", "RESILIENT", "Passes the approved One-Shock Test"),
        ("D", "Observed company health", "INSUFFICIENT_DATA", "Evidence window is not sufficient for a reliable assessment"),
        ("D", "Observed company health", "HEALTHY", "Commitments, capacity, debt and resilience remain in envelope"),
        ("D", "Observed company health", "LOAD_TIGHT", "One sustained signal approaches a boundary"),
        ("D", "Observed company health", "STRAINED", "Reserve is regularly consumed or debt grows"),
        ("D", "Observed company health", "UNSTABLE", "A typical disruption breaks commitments or operation is not durable"),
        ("E", "OWNER posture", "STABILITY", "Protect stability and reserve"),
        ("E", "OWNER posture", "GROWTH", "Higher load appetite within signed policy"),
        ("E", "OWNER posture", "PEAK", "Temporary peak posture with review date and exit condition"),
        ("E", "OWNER posture", "RECOVERY_POSTURE", "Deliberately repay debt and constrain intake"),
        ("F", "Fact or fact set", "SUFFICIENT", "Enough data for the named assessment"),
        ("F", "Fact or fact set", "INSUFFICIENT_DATA", "A material premise is missing"),
        ("F", "Fact or fact set", "STALE", "Data exceeded its validity for the named use"),
        ("F", "Fact or fact set", "CONFLICTING", "At least two credible sources conflict"),
    ]
    ws.append([])
    ws.append(["Axis", "Decision object", "State", "Normative meaning"])
    for row in rows:
        ws.append(row)
    _style_header(ws, 4, 1, 4)
    _style_body(ws, 5, 29, 1, 4)
    _table(ws, "A4:D29", "AxesStates")
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 32
    ws.column_dimensions["D"].width = 72
    for row in range(5, 30):
        ws.row_dimensions[row].height = 34
    ws.print_area = "A1:D29"


def _promise(wb: Workbook) -> None:
    ws = wb.create_sheet("02 Promise")
    ws.sheet_properties.tabColor = "C55A11"
    _sheet_base(ws, "Promise feasibility and execution gates", SOURCE, "K")
    headers = [
        "Case",
        "Technical feasibility",
        "Policy relation",
        "Override",
        "Critical data",
        "Safety",
        "Permission",
        "Consent",
        "Derived B",
        "Authorization",
        "Disposition / reason",
    ]
    rows = [
        ("Safe", "FEASIBLE", "WITHIN_ENVELOPE", "NOT_REQUIRED", "SUFFICIENT", "CLEAR", "VERIFIED", "NOT_REQUIRED", "SAFE", "ALLOWED", "EXECUTE"),
        ("Forced soft exception", "FEASIBLE", "SOFT_EXCEPTION", "APPROVED", "SUFFICIENT", "CLEAR", "VERIFIED", "NOT_REQUIRED", "FORCED", "ALLOWED", "EXECUTE_WITH_RECORDED_RISK"),
        ("Soft exception pending", "FEASIBLE", "SOFT_EXCEPTION", "PENDING", "SUFFICIENT", "CLEAR", "VERIFIED", "NOT_REQUIRED", "CONDITIONAL", "BLOCKED", "ABSTAIN · owner decision required"),
        ("Technical condition", "CONDITIONAL", "WITHIN_ENVELOPE", "NOT_REQUIRED", "SUFFICIENT", "CLEAR", "VERIFIED", "NOT_REQUIRED", "CONDITIONAL", "BLOCKED", "ABSTAIN · named condition outstanding"),
        ("Unknown feasibility", "UNKNOWN", "SOFT_EXCEPTION", "APPROVED", "SUFFICIENT", "CLEAR", "VERIFIED", "NOT_REQUIRED", "—", "BLOCKED", "ABSTAIN · unknown cannot become FORCED"),
        ("Technically infeasible", "INFEASIBLE", "SOFT_EXCEPTION", "APPROVED", "SUFFICIENT", "CLEAR", "VERIFIED", "NOT_REQUIRED", "INFEASIBLE", "BLOCKED", "ABSTAIN · override does not change truth"),
        ("Policy hard block", "FEASIBLE", "HARD_BLOCK", "APPROVED", "SUFFICIENT", "CLEAR", "VERIFIED", "NOT_REQUIRED", "INFEASIBLE", "BLOCKED", "FORBIDDEN"),
        ("Unverified method", "UNKNOWN", "WITHIN_ENVELOPE", "NOT_REQUIRED", "SUFFICIENT", "UNVERIFIED_CONCERN", "VERIFIED", "NOT_REQUIRED", "CONDITIONAL", "BLOCKED", "ABSTAIN · verify the method"),
        ("BHP hard block", "INFEASIBLE", "HARD_BLOCK", "APPROVED", "SUFFICIENT", "HARD_BLOCK", "VERIFIED", "NOT_REQUIRED", "INFEASIBLE", "BLOCKED", "FORBIDDEN"),
        ("Missing required permission", "FEASIBLE", "SOFT_EXCEPTION", "APPROVED", "SUFFICIENT", "CLEAR", "MISSING", "NOT_REQUIRED", "INFEASIBLE", "BLOCKED", "FORBIDDEN"),
        ("Unknown permission", "FEASIBLE", "SOFT_EXCEPTION", "APPROVED", "SUFFICIENT", "CLEAR", "UNKNOWN", "NOT_REQUIRED", "—", "BLOCKED", "ABSTAIN"),
        ("Consent declined", "FEASIBLE", "SOFT_EXCEPTION", "APPROVED", "SUFFICIENT", "CLEAR", "VERIFIED", "DECLINED", "CONDITIONAL", "BLOCKED", "ABSTAIN"),
        ("Critical data stale", "FEASIBLE", "SOFT_EXCEPTION", "APPROVED", "STALE", "CLEAR", "VERIFIED", "NOT_REQUIRED", "—", "BLOCKED", "ABSTAIN"),
        ("Critical data conflicting", "FEASIBLE", "SOFT_EXCEPTION", "APPROVED", "CONFLICTING", "CLEAR", "VERIFIED", "NOT_REQUIRED", "—", "BLOCKED", "ABSTAIN"),
    ]
    ws.append([])
    ws.append(headers)
    for row in rows:
        ws.append(row)
    _style_header(ws, 4, 1, 11)
    _style_body(ws, 5, 18, 1, 11)
    _table(ws, "A4:K18", "PromiseGateMatrix")
    for col, width in {
        "A": 28, "B": 22, "C": 21, "D": 16, "E": 18, "F": 23,
        "G": 18, "H": 18, "I": 16, "J": 16, "K": 40,
    }.items():
        ws.column_dimensions[col].width = width
    for row in range(5, 19):
        ws.row_dimensions[row].height = 40
    _status_colors(ws, "I", 5, 18)
    _status_colors(ws, "J", 5, 18)
    _status_colors(ws, "K", 5, 18)
    ws.print_area = "A1:K18"


def _health(wb: Workbook) -> None:
    ws = wb.create_sheet("03 Health Recovery")
    ws.sheet_properties.tabColor = "70AD47"
    _sheet_base(ws, "Health, recovery, and OWNER posture", SOURCE, "E")
    health = [
        ("INSUFFICIENT_DATA", "No reliable assessment; new company default"),
        ("HEALTHY", "Commitments, capacity, debt, and resilience within envelope"),
        ("LOAD_TIGHT", "One sustained boundary-approach signal"),
        ("STRAINED", "Reserve regularly consumed or debt grows"),
        ("UNSTABLE", "Typical disruption breaks commitments or operation is not durable"),
    ]
    ws.append([])
    ws.append(["Health assessment", "Meaning"])
    for row in health:
        ws.append(row)
    _style_header(ws, 4, 1, 2)
    _style_body(ws, 5, 9, 1, 2)
    _table(ws, "A4:B9", "HealthStates")

    _section_label(ws, 11, "Recovery control is not health", "E")
    ws.append(["Observed health", "Recovery signal", "OWNER decision", "Control state", "OWNER posture"])
    recovery = [
        ("UNSTABLE", "RECOVERY_NEEDED", "PENDING", "INACTIVE", "GROWTH"),
        ("UNSTABLE", "RECOVERY_NEEDED", "DECLINED", "INACTIVE", "GROWTH"),
        ("STRAINED", "RECOVERY_NEEDED", "ACTIVATED", "RECOVERY_REQUIRED", "RECOVERY_POSTURE"),
        ("HEALTHY", "NONE", "NOT_REQUESTED", "INACTIVE", "STABILITY"),
    ]
    for row in recovery:
        ws.append(row)
    _style_header(ws, 12, 1, 5)
    _style_body(ws, 13, 16, 1, 5)
    _table(ws, "A12:E16", "RecoverySeparation")
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 24
    ws.column_dimensions["D"].width = 28
    ws.column_dimensions["E"].width = 28
    for row in range(5, 17):
        ws.row_dimensions[row].height = 34
    ws.print_area = "A1:E16"


def _objections(wb: Workbook) -> None:
    ws = wb.create_sheet("04 Objections")
    ws.sheet_properties.tabColor = "FFC000"
    _sheet_base(ws, "Objections and consent", SOURCE, "D")
    objections = [
        ("PREFERENCE", "Record; no veto", "EXECUTE", "Content only; actor rank does not change result"),
        ("FACT_CORRECTION", "Reassess dependent facts", "ABSTAIN", "Resume only after fact evaluation"),
        ("PLAN_UNWORKABLE", "Reassess dependent plan", "ABSTAIN", "Named operational premise required"),
        ("AVAILABILITY_LIMIT", "Block dependent action", "ABSTAIN", "Required voluntary consent is absent"),
        ("SAFETY_OR_TECHNICAL_CONCERN", "Pause concrete method", "ABSTAIN", "Verify the method"),
        ("HARD_SAFETY_BLOCK", "Known hard boundary", "FORBIDDEN", "No override"),
    ]
    ws.append([])
    ws.append(["Objection type", "Gate effect", "Default outcome", "Normative note"])
    for row in objections:
        ws.append(row)
    _style_header(ws, 4, 1, 4)
    _style_body(ws, 5, 10, 1, 4)
    _table(ws, "A4:D10", "ObjectionMatrix")
    _status_colors(ws, "C", 5, 10)

    _section_label(ws, 12, "Consent gate", "D")
    consent = [
        ("NOT_REQUIRED", "Gate satisfied when consent is not legally or operationally required"),
        ("REQUESTED", "Pending; not consent"),
        ("FREELY_GIVEN", "Only state that satisfies a required consent gate"),
        ("DECLINED", "Dependent action blocked"),
        ("PRESSURED", "Invalid as voluntary consent; dependent action blocked"),
        ("ABSENT", "Silence is not consent; dependent action blocked"),
    ]
    ws.append(["Consent status", "Meaning"])
    for row in consent:
        ws.append(row)
    _style_header(ws, 13, 1, 2)
    _style_body(ws, 14, 19, 1, 2)
    _table(ws, "A13:B19", "ConsentStates")
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 56
    ws.column_dimensions["C"].width = 23
    ws.column_dimensions["D"].width = 56
    for row in range(5, 20):
        ws.row_dimensions[row].height = 36
    ws.print_area = "A1:D19"


def _voice_execution(wb: Workbook) -> None:
    ws = wb.create_sheet("05 Voice Execution")
    ws.sheet_properties.tabColor = "7030A0"
    _sheet_base(ws, "VOICE, NO_DECEPTION, and execution outcomes", SOURCE, "D")
    outcomes = [
        ("EXECUTE", "Full gate; no special recorded risk", "Concrete variant"),
        ("EXECUTE_WITH_RECORDED_RISK", "Legal and authorized; risk remains explicit", "Concrete variant"),
        ("ABSTAIN", "Insufficient basis, unmet condition, or required consent missing", "Dependent action only"),
        ("FORBIDDEN", "Known hard safety, legal, permission, or policy prohibition", "Forbidden variant"),
        ("HALT", "Critical control failure or threat to the entire flow", "Whole flow only"),
    ]
    ws.append([])
    ws.append(["Execution outcome", "Use", "Scope"])
    for row in outcomes:
        ws.append(row)
    _style_header(ws, 4, 1, 3)
    _style_body(ws, 5, 9, 1, 3)
    _table(ws, "A4:C9", "ExecutionOutcomes")
    _status_colors(ws, "A", 5, 9)

    _section_label(ws, 11, "Communication records", "D")
    records = [
        ("DRAFT_TEXT", "LLM output before deterministic validation", "Never system-approved"),
        ("APPROVED_TEXT", "Draft after deterministic validation and required approval", "Eligible for sending"),
        ("SENT_MESSAGE", "Message actually sent by the system", "Must derive from APPROVED_TEXT"),
        ("ACTUAL_CLAIM", "Credibly established human statement outside the system", "Never presented as system approval"),
    ]
    ws.append(["Record type", "Meaning", "Integrity rule"])
    for row in records:
        ws.append(row)
    _style_header(ws, 12, 1, 3)
    _style_body(ws, 13, 16, 1, 3)
    _table(ws, "A12:C16", "CommunicationRecords")

    _section_label(ws, 18, "VOICE validation", "D")
    voice = [
        ("VERIFIED", "FACT or weaker", "Evidence supports the assertion"),
        ("CONDITIONAL", "CONDITION", "Named material condition required"),
        ("UNKNOWN", "Explicit uncertainty or question", "No false fact or certainty"),
        ("FALSE", "No send", "Reject false claim"),
        ("GUARANTEE", "Only B=SAFE and F=SUFFICIENT", "FORCED never authorizes false certainty"),
    ]
    ws.append(["Evidence / rendering", "Allowed communication", "Rule"])
    for row in voice:
        ws.append(row)
    _style_header(ws, 19, 1, 3)
    _style_body(ws, 20, 24, 1, 3)
    _table(ws, "A19:C24", "VoiceValidation")
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 62
    ws.column_dimensions["C"].width = 46
    ws.column_dimensions["D"].width = 3
    for row in range(5, 25):
        ws.row_dimensions[row].height = 38
    ws.print_area = "A1:D24"


def _regression(wb: Workbook) -> None:
    ws = wb.create_sheet("06 Regression 1017")
    ws.sheet_properties.tabColor = "C00000"
    _sheet_base(
        ws,
        "DOM / KLINIKA / LOFT — Wednesday 10:17",
        "Blocking regression contract · later outcomes never rewrite the 10:17 assessment",
        "C",
    )
    cases = [
        ("T-OBJ-01", "Crew move, promise, and method are different objects", "Separate assessments; no forced state sharing"),
        ("T-SCOPE-01", "Locator failure is independent of DOM protection", "Does not block T1"),
        ("T-PREF-01", "Marek states a preference without a new fact", "No veto; T1 allowed with recorded reserve risk"),
        ("T-AVAIL-01", "Anna declines availability requiring consent", "T2 = ABSTAIN"),
        ("T-METHOD-01", "Material technical condition is unverified", "Concrete T3 method = ABSTAIN"),
        ("T-BHP-01", "Method violates a confirmed prohibition", "T3 = FORBIDDEN"),
        ("T-VOICE-01", "‘Definitely tomorrow’ with critical F ≠ SUFFICIENT", "No approved text and no send"),
        ("T-CLAIM-01", "OWNER makes the promise outside the system", "A = PROMISE; ACTUAL_CLAIM; B/F unchanged"),
        ("T-FORCED-01", "INFEASIBLE + OWNER override", "Remains INFEASIBLE; never FORCED"),
        ("T-FORCED-02", "FEASIBLE + SOFT_EXCEPTION + APPROVED", "B = FORCED; complete trail required"),
        ("T-REC-01", "D = UNSTABLE; OWNER declines recovery", "D unchanged; decision = DECLINED"),
        ("T-HYS-01", "One single-day signal", "No health transition"),
        ("T-OUTCOME-01", "Later success or failure", "Does not rewrite the 10:17 assessment"),
        ("T-ROLE-01", "Speaker and OWNER roles are reversed", "Same evidence gets the same content classification"),
        ("T-DECEPTION-01", "Goal would require lying or simulated consent", "Reject message; offer honest alternative"),
    ]
    ws.append([])
    ws.append(["Test ID", "Input at 10:17", "Expected result"])
    for row in cases:
        ws.append(row)
    _style_header(ws, 4, 1, 3)
    _style_body(ws, 5, 19, 1, 3)
    _table(ws, "A4:C19", "RegressionContract")
    _status_colors(ws, "C", 5, 19)
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 68
    ws.column_dimensions["C"].width = 68
    for row in range(5, 20):
        ws.row_dimensions[row].height = 44
    ws.print_area = "A1:C19"


def build(output: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.title = "WERKcrew State & Transition Matrix v1.1"
    wb.properties.subject = "Decision Semantics Freeze publication matrix"
    wb.properties.creator = "WERKcrew"
    wb.properties.keywords = "WERKcrew, decision semantics, governance, matrix"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    _overview(wb)
    _axes(wb)
    _promise(wb)
    _health(wb)
    _objections(wb)
    _voice_execution(wb)
    _regression(wb)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)

    # Reopen immediately so truncated/invalid OOXML cannot pass silently.
    checked = load_workbook(output, data_only=False, read_only=False)
    expected = [
        "00 Release",
        "01 Axes",
        "02 Promise",
        "03 Health Recovery",
        "04 Objections",
        "05 Voice Execution",
        "06 Regression 1017",
    ]
    if checked.sheetnames != expected:
        raise RuntimeError(f"Unexpected worksheet topology: {checked.sheetnames!r}")
    checked.close()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--output",
        type=Path,
        default=Path(
            "docs/governance/releases/v1.1/"
            "WERKcrew_State_Transition_Matrix_v1.1.xlsx"
        ),
    )
    args = parser.parse_args()
    build(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
